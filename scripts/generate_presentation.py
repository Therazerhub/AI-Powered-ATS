#!/usr/bin/env python3
"""Generate the TalentIQ / AI-Powered ATS project presentation."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "AI_POWERED_ATS_PRESENTATION.pptx"

# Palette mirrors static/css/app.css.
DARK = "102C28"
INK = "162321"
GREEN = "155E52"
GREEN2 = "237D6D"
MINT = "5BD7B6"
PALE_MINT = "E7F4F0"
CREAM = "F5F7F4"
WHITE = "FFFFFF"
MUTED = "667572"
LINE = "DFE6E2"
AMBER = "D69338"
RED = "B85450"
BLUE = "4263EB"
SOFT_BLUE = "E9EDFF"

SW = Inches(13.333)
SH = Inches(7.5)


def rgb(hex_color):
    return RGBColor.from_string(hex_color)


def add_shape(slide, shape_type, x, y, w, h, fill, line=None, radius=False):
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line or fill)
    if radius:
        shape.adjustments[0] = 0.12
    return shape


def add_line(slide, x1, y1, x2, y2, color=LINE, width=1.5, arrow=False):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    line.line.color.rgb = rgb(color)
    line.line.width = Pt(width)
    if arrow:
        line.line.end_arrowhead = True
    return line


def add_text(slide, text, x, y, w, h, size=18, color=INK, bold=False,
             font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
             margin=0.03, fit=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    if fit:
        tf.fit_text(font_family=font, max_size=size)
    return box


def add_rich_text(slide, runs, x, y, w, h, size=18, color=INK,
                  align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.02)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    for item in runs:
        run = p.add_run()
        run.text = item[0]
        run.font.name = "Aptos"
        run.font.size = Pt(item[1] if len(item) > 1 else size)
        run.font.bold = item[2] if len(item) > 2 else False
        run.font.color.rgb = rgb(item[3] if len(item) > 3 else color)
    return box


def add_bullets(slide, items, x, y, w, h, size=15, color=INK,
                bullet_color=GREEN2, spacing=8):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.03)
    tf.margin_right = Inches(0.03)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(spacing)
        p.line_spacing = 1.08
        r1 = p.add_run()
        r1.text = "●  "
        r1.font.name = "Aptos"
        r1.font.size = Pt(max(8, size - 5))
        r1.font.color.rgb = rgb(bullet_color)
        r2 = p.add_run()
        if isinstance(item, tuple):
            lead, body = item
            r2.text = lead
            r2.font.bold = True
            r2.font.size = Pt(size)
            r2.font.color.rgb = rgb(color)
            r3 = p.add_run()
            r3.text = body
            r3.font.size = Pt(size)
            r3.font.color.rgb = rgb(color)
            r3.font.name = "Aptos"
        else:
            r2.text = item
            r2.font.size = Pt(size)
            r2.font.color.rgb = rgb(color)
        r2.font.name = "Aptos"
    return box


def background(slide, color=CREAM):
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = rgb(color)


def brand(slide, light=False):
    fill = MINT if light else DARK
    ink = DARK if light else WHITE
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 0.55, 0.34, 0.34, 0.34, fill, fill, True)
    add_text(slide, "✦", 0.55, 0.335, 0.34, 0.34, 15, ink, True,
             align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, "TalentIQ", 0.98, 0.34, 1.4, 0.32, 14,
             WHITE if light else DARK, True, valign=MSO_ANCHOR.MIDDLE)


def footer(slide, num, dark=False):
    color = "9CB7B0" if dark else MUTED
    add_text(slide, "AI-POWERED ATS  •  PROJECT OVERVIEW", 0.6, 7.10, 4.0, 0.2, 8, color, True)
    add_text(slide, f"{num:02d} / 10", 12.1, 7.08, 0.63, 0.22, 9, color, True, align=PP_ALIGN.RIGHT)


def section_title(slide, kicker, title, subtitle, num, dark=False):
    brand(slide, light=dark)
    tc = WHITE if dark else INK
    sc = "BFD2CD" if dark else MUTED
    add_text(slide, kicker.upper(), 0.62, 0.92, 3.5, 0.25, 10, MINT if dark else GREEN2, True)
    add_text(slide, title, 0.62, 1.16, 11.9, 0.65, 29, tc, True)
    if subtitle:
        add_text(slide, subtitle, 0.64, 1.82, 11.6, 0.46, 13, sc)
    footer(slide, num, dark)


def pill(slide, text, x, y, w, fill=PALE_MINT, color=GREEN, border=None, size=10):
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, 0.34, fill, border or fill, True)
    add_text(slide, text, x, y + 0.01, w, 0.30, size, color, True,
             align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)


def card(slide, x, y, w, h, title, body=None, accent=GREEN2, number=None,
         fill=WHITE, title_size=16):
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h, fill, LINE, True)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, x, y, 0.07, h, accent, accent)
    if number is not None:
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x + 0.28, y + 0.28, 0.54, 0.54, accent, accent)
        add_text(slide, str(number), x + 0.28, y + 0.28, 0.54, 0.54, 15, WHITE, True,
                 align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        tx = x + 0.98
        tw = w - 1.25
    else:
        tx = x + 0.3
        tw = w - 0.55
    add_text(slide, title, tx, y + 0.25, tw, 0.38, title_size, INK, True)
    if body:
        add_text(slide, body, tx, y + 0.72, tw, h - 0.88, 11.5, MUTED)


def entity(slide, name, fields, x, y, w, h, accent=GREEN2):
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h, WHITE, LINE, True)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, x, y, w, 0.42, accent, accent)
    add_text(slide, name, x + 0.12, y + 0.05, w - 0.24, 0.3, 12, WHITE, True, valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, "\n".join(fields), x + 0.16, y + 0.54, w - 0.32, h - 0.62, 9.5, INK)


def build():
    prs = Presentation()
    prs.slide_width = SW
    prs.slide_height = SH
    blank = prs.slide_layouts[6]
    prs.core_properties.title = "AI-Powered ATS — TalentIQ"
    prs.core_properties.subject = "Django recruitment workflow with AI-assisted candidate analysis"
    prs.core_properties.author = "TalentIQ project team"
    prs.core_properties.keywords = "ATS, Django, OpenRouter, recruitment, AI"

    # 1 — Title
    slide = prs.slides.add_slide(blank)
    background(slide, DARK)
    # Decorative geometry inspired by the web hero.
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, 9.65, -0.55, 4.5, 4.5, GREEN, GREEN)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, 10.55, 0.12, 3.2, 3.2, GREEN2, GREEN2)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, 11.25, 0.72, 1.85, 1.85, MINT, MINT)
    brand(slide, light=True)
    pill(slide, "DJANGO  •  SQLITE  •  OPENROUTER", 0.68, 1.44, 3.25, fill="1C463E", color="A8E8D7", size=9)
    add_rich_text(slide, [
        ("AI-Powered ", 39, True, WHITE),
        ("Applicant\nTracking System", 39, True, MINT),
    ], 0.68, 2.0, 7.4, 1.78)
    add_text(slide, "A skills-first recruitment workflow for candidates and recruiters — with transparent, advisory AI analysis.",
             0.72, 4.02, 6.65, 0.85, 17, "D4E5E0")
    add_line(slide, 0.72, 5.22, 6.55, 5.22, "356259", 1)
    add_text(slide, "PROJECT PRESENTATION", 0.72, 5.46, 2.0, 0.25, 10, MINT, True)
    add_text(slide, "TalentIQ", 0.72, 5.78, 2.8, 0.45, 22, WHITE, True)
    # Mini dashboard card
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 8.25, 2.33, 3.55, 3.65, "173B35", "356259", True)
    add_text(slide, "CANDIDATE MATCH", 8.62, 2.66, 1.75, 0.26, 10, "BFD2CD", True)
    pill(slide, "LIVE AI", 10.58, 2.61, 0.82, fill=GREEN2, color=WHITE, size=8)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, 9.22, 3.16, 1.55, 1.55, "21463E", MINT)
    add_text(slide, "88%", 9.22, 3.37, 1.55, 0.52, 28, WHITE, True, align=PP_ALIGN.CENTER)
    add_text(slide, "advisory fit", 9.22, 4.03, 1.55, 0.24, 9, "BFD2CD", align=PP_ALIGN.CENTER)
    add_text(slide, "Python", 8.62, 5.02, 1.1, 0.25, 11, WHITE, True)
    add_text(slide, "MATCH", 10.53, 5.02, 0.72, 0.25, 9, MINT, True, align=PP_ALIGN.RIGHT)
    add_line(slide, 8.62, 5.38, 11.38, 5.38, "356259", 1)
    add_text(slide, "Human decision required", 8.62, 5.55, 2.76, 0.24, 9, "BFD2CD")
    footer(slide, 1, True)

    # 2 — Problem
    slide = prs.slides.add_slide(blank)
    background(slide)
    section_title(slide, "01  /  Context", "Recruitment breaks when evidence gets fragmented",
                  "The project addresses workflow friction—not a lack of candidate data.", 2)
    card(slide, 0.65, 2.55, 3.7, 2.47, "Scattered qualification data",
         "Profiles, skills, projects, education and PDF resumes are difficult to compare consistently.", AMBER, 1)
    card(slide, 4.82, 2.55, 3.7, 2.47, "Manual pipeline overhead",
         "Recruiters need one place to manage jobs, applications, status changes and interviews.", RED, 2)
    card(slide, 8.99, 2.55, 3.7, 2.47, "Opaque screening risk",
         "A score without evidence or human review can turn assistance into an unaccountable decision.", BLUE, 3)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 1.28, 5.55, 10.78, 0.86, DARK, DARK, True)
    add_text(slide, "Design question", 1.62, 5.78, 1.4, 0.3, 11, MINT, True)
    add_text(slide, "How can one system organize the workflow, surface job-relevant evidence, and keep recruiters accountable?",
             3.1, 5.69, 8.52, 0.48, 15, WHITE, True, valign=MSO_ANCHOR.MIDDLE)

    # 3 — Solution
    slide = prs.slides.add_slide(blank)
    background(slide, DARK)
    section_title(slide, "02  /  Product", "One workflow from profile to hiring decision",
                  "TalentIQ connects candidate records, job requirements and recruiter actions.", 3, True)
    steps = [
        ("01", "Build profile", "Skills, education, experience, projects, certifications"),
        ("02", "Apply with PDF", "Active resume and optional cover letter"),
        ("03", "Compare evidence", "AI-assisted match, gaps, strengths and concerns"),
        ("04", "Move the pipeline", "Review, shortlist, interview, select or reject"),
    ]
    xs = [0.7, 3.82, 6.94, 10.06]
    for x, (n, title, body) in zip(xs, steps):
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, 2.72, 2.57, 2.55, "173B35", "356259", True)
        add_text(slide, n, x + 0.25, 2.96, 0.48, 0.35, 13, MINT, True)
        add_text(slide, title, x + 0.25, 3.48, 2.05, 0.52, 17, WHITE, True)
        add_text(slide, body, x + 0.25, 4.16, 2.05, 0.76, 11, "BFD2CD")
    for x in [3.36, 6.48, 9.60]:
        add_text(slide, "→", x, 3.73, 0.38, 0.38, 19, MINT, True, align=PP_ALIGN.CENTER)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 2.17, 5.73, 8.98, 0.59, "1C463E", "356259", True)
    add_text(slide, "Shared system of record: Django models + SQLite + role-aware views", 2.17, 5.86, 8.98, 0.3,
             12, "D4E5E0", True, align=PP_ALIGN.CENTER)

    # 4 — Roles & features
    slide = prs.slides.add_slide(blank)
    background(slide)
    section_title(slide, "03  /  Experience", "Two role-aware journeys, one recruitment record",
                  "Views route authenticated users to candidate or recruiter dashboards.", 4)
    # Candidate panel
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 0.67, 2.46, 5.78, 3.78, WHITE, LINE, True)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 0.91, 2.73, 0.72, 0.72, PALE_MINT, PALE_MINT, True)
    add_text(slide, "C", 0.91, 2.73, 0.72, 0.72, 20, GREEN, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, "Candidate workspace", 1.84, 2.75, 3.3, 0.35, 19, INK, True)
    add_text(slide, "SELF-SERVICE", 1.85, 3.16, 1.45, 0.24, 9, GREEN2, True)
    add_bullets(slide, [
        ("Profile & skills — ", "capture education, experience, projects and certifications"),
        ("Resume management — ", "upload PDF and select an active resume"),
        ("Job discovery — ", "search open roles and inspect requirements"),
        ("Applications — ", "apply once per job and track status"),
    ], 1.0, 3.65, 5.0, 2.2, 12.5)
    # Recruiter panel
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 6.88, 2.46, 5.78, 3.78, DARK, DARK, True)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 7.13, 2.73, 0.72, 0.72, GREEN2, GREEN2, True)
    add_text(slide, "R", 7.13, 2.73, 0.72, 0.72, 20, WHITE, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, "Recruiter workspace", 8.06, 2.75, 3.3, 0.35, 19, WHITE, True)
    add_text(slide, "DECISION SUPPORT", 8.07, 3.16, 1.75, 0.24, 9, MINT, True)
    add_bullets(slide, [
        ("Jobs — ", "create, edit, close and inspect applicant counts"),
        ("Candidate search — ", "find people by name, email or skill"),
        ("Pipeline — ", "update six statuses and add recruiter notes"),
        ("Interviews & AI — ", "schedule details and trigger analysis"),
    ], 7.23, 3.65, 5.0, 2.2, 12.5, "DDEBE7", MINT)

    # 5 — Django MVT
    slide = prs.slides.add_slide(blank)
    background(slide)
    section_title(slide, "04  /  Architecture", "Django MVT keeps responsibilities explicit",
                  "The recruitment app turns HTTP requests into validated, persisted workflows.", 5)
    layers = [
        ("T", "Templates", "Bootstrap UI\n17 HTML templates", BLUE, SOFT_BLUE),
        ("V", "Views + Forms", "Role decorators\nvalidation + messages", GREEN2, PALE_MINT),
        ("M", "Models", "9 domain models\nORM relationships", AMBER, "FFF4E3"),
    ]
    for i, (letter, title, body, accent, pale) in enumerate(layers):
        x = 0.75 + i * 3.47
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, 2.55, 2.95, 2.5, WHITE, LINE, True)
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x + 0.24, 2.84, 0.68, 0.68, pale, pale)
        add_text(slide, letter, x + 0.24, 2.84, 0.68, 0.68, 20, accent, True,
                 align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        add_text(slide, title, x + 1.07, 2.87, 1.58, 0.34, 17, INK, True)
        add_text(slide, body, x + 0.25, 3.77, 2.45, 0.7, 13, MUTED, align=PP_ALIGN.CENTER)
    # AI service sidecar
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 11.14, 2.55, 1.45, 2.5, DARK, DARK, True)
    add_text(slide, "AI", 11.45, 2.86, 0.83, 0.56, 21, MINT, True, align=PP_ALIGN.CENTER)
    add_text(slide, "services.py", 11.33, 3.57, 1.08, 0.27, 11, WHITE, True, align=PP_ALIGN.CENTER)
    add_text(slide, "PDF text\nOpenRouter\nJSON parse", 11.31, 4.03, 1.1, 0.74, 10, "BFD2CD", align=PP_ALIGN.CENTER)
    for x in [3.77, 7.24, 10.70]:
        add_text(slide, "→", x, 3.52, 0.3, 0.4, 20, GREEN2, True, align=PP_ALIGN.CENTER)
    # Request lifecycle
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 1.18, 5.58, 10.95, 0.61, "EAF1EE", "D5E1DD", True)
    add_text(slide, "URL  →  view/decorator  →  form validation  →  ORM/service  →  template or JSON response",
             1.18, 5.73, 10.95, 0.3, 12.5, GREEN, True, align=PP_ALIGN.CENTER)
    pill(slide, "SQLite", 5.82, 6.42, 1.18, fill=DARK, color=WHITE, size=10)

    # 6 — Database relationships
    slide = prs.slides.add_slide(blank)
    background(slide, DARK)
    section_title(slide, "05  /  Data model", "Relationships preserve the recruitment story",
                  "Foreign keys, one-to-one records and shared skills connect every decision to evidence.", 6, True)
    # Connectors first so cards sit above them.
    add_line(slide, 2.45, 3.28, 4.12, 3.28, "5F857B", 1.3)
    add_line(slide, 5.82, 3.28, 7.55, 3.28, "5F857B", 1.3)
    add_line(slide, 9.23, 3.28, 10.83, 3.28, "5F857B", 1.3)
    add_line(slide, 8.39, 4.17, 8.39, 5.18, "5F857B", 1.3)
    add_line(slide, 8.39, 5.18, 4.64, 5.18, "5F857B", 1.3)
    add_line(slide, 8.39, 5.18, 8.39, 5.46, "5F857B", 1.3)
    add_line(slide, 8.39, 5.18, 11.47, 5.18, "5F857B", 1.3)
    entity(slide, "User", ["Django auth", "1:1 role profile"], 0.68, 2.48, 1.78, 1.48, GREEN2)
    entity(slide, "CandidateProfile", ["skills M:N", "education • experience"], 4.08, 2.32, 1.92, 1.78, GREEN2)
    entity(slide, "Application", ["status • cover letter", "AI score + JSON"], 7.48, 2.32, 1.92, 1.78, MINT)
    entity(slide, "Job", ["required/preferred skills", "deadline • status"], 10.82, 2.32, 1.84, 1.78, AMBER)
    entity(slide, "Resume", ["PDF • extracted text", "active flag"], 3.68, 5.30, 1.92, 1.20, BLUE)
    entity(slide, "AIAnalysis", ["evidence lists • model"], 7.48, 5.30, 1.92, 1.20, GREEN2)
    entity(slide, "Interview", ["time • mode • link"], 10.54, 5.30, 1.86, 1.20, AMBER)
    entity(slide, "RecruiterProfile", ["title • department"], 0.68, 5.30, 2.12, 1.20, GREEN2)
    pill(slide, "1 per candidate + job", 6.66, 4.49, 2.37, fill="1C463E", color="BFD2CD", border="356259", size=9)
    pill(slide, "Skill M:N", 9.78, 4.49, 1.38, fill="1C463E", color=MINT, border="356259", size=9)

    # 7 — AI flow
    slide = prs.slides.add_slide(blank)
    background(slide)
    section_title(slide, "06  /  AI integration", "The API flow is constrained, parsed and auditable",
                  "A recruiter triggers analysis for one application; the result is advisory evidence.", 7)
    flow = [
        ("1", "Extract", "pypdf reads active PDF\n(up to 30k chars)"),
        ("2", "Minimize", "Job qualifications +\nanonymized candidate payload"),
        ("3", "Request", "OpenRouter chat API\ntemperature 0.1"),
        ("4", "Validate", "Required JSON keys\nscore clamped 0–100"),
        ("5", "Persist", "Application JSON +\nAIAnalysis record"),
        ("6", "Review", "Recruiter reads evidence\nand decides"),
    ]
    for i, (n, title, body) in enumerate(flow):
        x = 0.56 + i * 2.12
        accent = GREEN2 if i < 5 else DARK
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, 2.73, 1.72, 2.44, WHITE, LINE, True)
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x + 0.56, 2.98, 0.60, 0.60, accent, accent)
        add_text(slide, n, x + 0.56, 2.98, 0.60, 0.60, 15, WHITE, True,
                 align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        add_text(slide, title, x + 0.15, 3.79, 1.42, 0.32, 14.5, INK, True, align=PP_ALIGN.CENTER)
        add_text(slide, body, x + 0.13, 4.28, 1.46, 0.64, 9.7, MUTED, align=PP_ALIGN.CENTER)
        if i < 5:
            add_text(slide, "→", x + 1.78, 3.62, 0.30, 0.34, 17, GREEN2, True, align=PP_ALIGN.CENTER)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 1.05, 5.65, 11.25, 0.74, DARK, DARK, True)
    add_text(slide, "RETURNED EVIDENCE", 1.35, 5.88, 1.55, 0.25, 9, MINT, True)
    add_text(slide, "match score  •  matching & missing skills  •  relevant experience/projects  •  strengths  •  concerns  •  explanation",
             3.02, 5.80, 8.88, 0.40, 11.2, WHITE, True, valign=MSO_ANCHOR.MIDDLE)

    # 8 — Ethics/security
    slide = prs.slides.add_slide(blank)
    background(slide)
    section_title(slide, "07  /  Governance", "Responsible assistance needs technical guardrails",
                  "The code implements baseline controls; production hardening remains an explicit step.", 8)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 0.7, 2.47, 5.88, 3.94, WHITE, LINE, True)
    pill(slide, "IMPLEMENTED IN PROJECT", 1.01, 2.79, 2.02, fill=PALE_MINT, color=GREEN, size=9)
    add_bullets(slide, [
        ("Skills-first prompt — ", "protected traits are explicitly excluded"),
        ("Human authority — ", "score is advisory, not a hiring decision"),
        ("Access control — ", "login plus candidate/recruiter decorators"),
        ("Input constraints — ", "PDF-only resume, 5 MB form limit"),
        ("Integrity — ", "unique application per candidate and job"),
        ("Web protections — ", "Django CSRF and security middleware"),
    ], 1.02, 3.43, 5.12, 2.55, 11.4)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 6.89, 2.47, 5.74, 3.94, DARK, DARK, True)
    pill(slide, "BEFORE PRODUCTION", 7.21, 2.79, 1.79, fill="1C463E", color=MINT, border="356259", size=9)
    add_bullets(slide, [
        ("Environment — ", "disable DEBUG; set a strong SECRET_KEY"),
        ("Authentication — ", "enable password validators and review staff provisioning"),
        ("Transport — ", "HTTPS, secure cookies and trusted origins"),
        ("Data lifecycle — ", "retention, deletion and encrypted backups"),
        ("Oversight — ", "audit logs, bias checks and recruiter training"),
        ("Resilience — ", "rate limits, timeouts and provider failure handling"),
    ], 7.23, 3.43, 4.98, 2.55, 11.4, "DDEBE7", MINT)

    # 9 — Demo sequence
    slide = prs.slides.add_slide(blank)
    background(slide, DARK)
    section_title(slide, "08  /  Demonstration", "A seven-minute proof of the complete workflow",
                  "Follow one candidate and one role from setup to recruiter decision support.", 9, True)
    demo = [
        ("01", "Recruiter creates role", "Requirements, skills, deadline"),
        ("02", "Candidate registers", "Profile, skills and PDF resume"),
        ("03", "Candidate applies", "Resume + cover letter; duplicate blocked"),
        ("04", "Recruiter opens queue", "Candidates ranked by available AI score"),
        ("05", "Run live analysis", "Inspect fit, gaps and explanation"),
        ("06", "Advance pipeline", "Shortlist and schedule interview"),
    ]
    for i, (n, title, body) in enumerate(demo):
        row, col = divmod(i, 3)
        x = 0.73 + col * 4.2
        y = 2.55 + row * 1.78
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, 3.74, 1.38, "173B35", "356259", True)
        add_text(slide, n, x + 0.24, y + 0.26, 0.48, 0.3, 11, MINT, True)
        add_text(slide, title, x + 0.86, y + 0.22, 2.55, 0.34, 15, WHITE, True)
        add_text(slide, body, x + 0.86, y + 0.70, 2.53, 0.45, 10.4, "BFD2CD")
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 3.02, 6.22, 7.28, 0.52, MINT, MINT, True)
    add_text(slide, "End state: evidence visible, status recorded, interview linked to the application",
             3.02, 6.34, 7.28, 0.28, 11, DARK, True, align=PP_ALIGN.CENTER)

    # 10 — Conclusion
    slide = prs.slides.add_slide(blank)
    background(slide)
    section_title(slide, "09  /  Close", "TalentIQ makes recruitment evidence easier to act on",
                  "The project combines practical workflow design with bounded AI assistance.", 10)
    outcomes = [
        ("ORGANIZE", "A shared record for jobs, candidates, resumes and applications", GREEN2),
        ("EXPLAIN", "Structured AI output exposes evidence—not just a score", BLUE),
        ("DECIDE", "Role-aware workflow keeps the recruiter responsible", AMBER),
    ]
    for i, (label, text, accent) in enumerate(outcomes):
        x = 0.73 + i * 4.17
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, 2.62, 3.70, 2.12, WHITE, LINE, True)
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, x, 2.62, 3.70, 0.09, accent, accent)
        add_text(slide, label, x + 0.28, 2.97, 2.2, 0.25, 10, accent, True)
        add_text(slide, text, x + 0.28, 3.47, 3.03, 0.82, 15, INK, True)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 0.73, 5.20, 11.78, 1.04, DARK, DARK, True)
    add_text(slide, "NEXT", 1.07, 5.56, 0.62, 0.27, 10, MINT, True)
    add_text(slide, "Harden deployment", 1.92, 5.47, 2.05, 0.38, 14, WHITE, True)
    add_text(slide, "→", 4.06, 5.44, 0.42, 0.42, 18, MINT, True, align=PP_ALIGN.CENTER)
    add_text(slide, "Add auditability", 4.66, 5.47, 1.84, 0.38, 14, WHITE, True)
    add_text(slide, "→", 6.66, 5.44, 0.42, 0.42, 18, MINT, True, align=PP_ALIGN.CENTER)
    add_text(slide, "Measure fairness & quality", 7.27, 5.47, 2.55, 0.38, 14, WHITE, True)
    add_text(slide, "→", 9.96, 5.44, 0.42, 0.42, 18, MINT, True, align=PP_ALIGN.CENTER)
    add_text(slide, "Pilot", 10.61, 5.47, 0.9, 0.38, 14, WHITE, True)
    add_text(slide, "Questions?", 4.82, 6.56, 3.7, 0.4, 20, GREEN, True, align=PP_ALIGN.CENTER)

    prs.save(OUTPUT)
    print(f"Created {OUTPUT}")
    print(f"Slides: {len(prs.slides)}")


if __name__ == "__main__":
    build()
